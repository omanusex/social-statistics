"""Renders chapter content (a list of blocks) into a Word document with python-docx."""
import os
import re
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Emu
from PIL import Image

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
FIG = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "figures"))
BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"


# ---- block constructors used by the chapter files ----------------------------
def H1(t): return ("h1", t)
def H2(t): return ("h2", t)
def H3(t): return ("h3", t)
def P(t): return ("p", t)
def BUL(items): return ("bul", items)
def NUM(items): return ("num", items)
def TABLE(caption, header, rows, widths=None, size=10): return ("table", caption, header, rows, widths, size)
def FIG_(path, caption, width_cm=16, max_h_cm=21): return ("fig", path, caption, width_cm, max_h_cm)
def CODE(caption, file, start, end=None, lines=None, skip_blank_end=True): return ("code", caption, file, start, end, lines)
def LANDSCAPE(): return ("orient", "landscape")
def PORTRAIT(): return ("orient", "portrait")
def BREAK(): return ("break",)
def NOTE(t): return ("note", t)
def CHAPTER(n): return ("chapter", n)
def REF(t): return ("ref", t)


def extract(file, start, end=None, lines=None):
    """Exact lines from a source file: from the first line matching `start`
    to the first later line matching `end` (inclusive), or `lines` lines."""
    src = open(os.path.join(ROOT, file), encoding="utf-8").read().splitlines()
    i = next(n for n, l in enumerate(src) if re.search(start, l))
    if lines:
        j = i + lines
    else:
        j = next(n for n in range(i + 1, len(src)) if re.search(end, src[n])) + 1
    chunk = src[i:j]
    # remove common indentation
    ind = min((len(l) - len(l.lstrip()) for l in chunk if l.strip()), default=0)
    return [l[ind:] for l in chunk], i + 1, j


class Builder:
    def __init__(self, chapter_no):
        self.doc = Document()
        self.ch = chapter_no
        self.fig_no = 0
        self.tab_no = 0
        self.lst_no = 0
        self._styles()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(2.54))
        self._footer(sec)
        self.first_block = True

    # ---- styles --------------------------------------------------------------
    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(12)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        pf = normal.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(6)
        for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
            h = st[name]
            h.font.name = BODY_FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = RGBColor(0, 0, 0)
            rpr = h.element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.append(rfonts)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rfonts.set(qn(a), BODY_FONT)
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rfonts.attrib.pop(qn(a), None)
            h.paragraph_format.space_before = Pt(18 if name == "Heading 1" else 12)
            h.paragraph_format.space_after = Pt(8 if name == "Heading 1" else 6)
            h.paragraph_format.line_spacing = 1.15
            h.paragraph_format.keep_with_next = True
        cap = st["Caption"]
        cap.font.name = BODY_FONT
        cap.font.size = Pt(10)
        cap.font.bold = True
        cap.font.italic = False
        cap.font.color.rgb = RGBColor(0, 0, 0)

    def _footer(self, section):
        section.footer.is_linked_to_previous = True
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if not p.runs:
            run = p.add_run()
            for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
                if kind:
                    el = OxmlElement("w:fldChar")
                    el.set(qn("w:fldCharType"), kind)
                else:
                    el = OxmlElement("w:instrText")
                    el.set(qn("xml:space"), "preserve")
                    el.text = text
                run._r.append(el)
            run.font.size = Pt(10)

    # ---- inline text: `code` and **bold** ------------------------------------
    def _runs(self, par, text, size=None):
        for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
            if not part:
                continue
            if part.startswith("`"):
                r = par.add_run(part[1:-1])
                r.font.name = CODE_FONT
                r._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
                r.font.size = Pt((size or 12) - 1.5)
            elif part.startswith("**"):
                r = par.add_run(part[2:-2])
                r.bold = True
                if size:
                    r.font.size = Pt(size)
            else:
                r = par.add_run(part)
                if size:
                    r.font.size = Pt(size)

    def _caption(self, label, no, text, before=True):
        p = self.doc.add_paragraph(style="Caption")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4 if not before else 8)
        p.paragraph_format.space_after = Pt(10 if not before else 4)
        p.paragraph_format.keep_with_next = before
        r = p.add_run(f"{label} {self.ch}.{no}: ")
        r.bold = True
        r2 = p.add_run(text)
        r2.bold = False
        return p

    # ---- orientation -----------------------------------------------------------
    def _orient(self, kind):
        current = self.doc.sections[-1]
        want_land = kind == "landscape"
        if (current.orientation == WD_ORIENT.LANDSCAPE) == want_land:
            return
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.orientation = WD_ORIENT.LANDSCAPE if want_land else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = (Cm(29.7), Cm(21)) if want_land else (Cm(21), Cm(29.7))
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Cm(2.0 if want_land else 2.54))
        self._footer(sec)

    def _usable(self):
        s = self.doc.sections[-1]
        return (s.page_width - s.left_margin - s.right_margin), (s.page_height - s.top_margin - s.bottom_margin)

    # ---- rendering ---------------------------------------------------------------
    def add(self, blocks):
        for b in blocks:
            getattr(self, "_b_" + b[0])(*b[1:])

    def _b_chapter(self, n):
        self.ch, self.fig_no, self.tab_no, self.lst_no = n, 0, 0, 0

    def _b_h1(self, t):
        if not self.first_block:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self.first_block = False
        p = self.doc.add_heading(t, level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _b_h2(self, t):
        self.doc.add_heading(t, level=2)

    def _b_h3(self, t):
        self.doc.add_heading(t, level=3)

    def _b_p(self, t):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self._runs(p, t)

    def _b_ref(self, t):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.first_line_indent = Cm(-1)
        self._runs(p, t)

    def _b_note(self, t):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        r = p.add_run(t)
        r.italic = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    def _list(self, items, style):
        for it in items:
            p = self.doc.add_paragraph(style=style)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(3)
            self._runs(p, it)

    def _b_bul(self, items):
        self._list(items, "List Bullet")

    def _b_num(self, items):
        self._list(items, "List Number")

    def _b_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _b_orient(self, kind):
        self._orient(kind)

    def _b_table(self, caption, header, rows, widths, size):
        self.tab_no += 1
        self._caption("Table", self.tab_no, caption, before=True)
        uw, _ = self._usable()
        n = len(header)
        widths = widths or [1] * n
        tw = sum(widths)
        col_w = [int(uw * w / tw) for w in widths]
        t = self.doc.add_table(rows=1, cols=n)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            par = c.paragraphs[0]
            r = par.add_run(h)
            r.bold = True
            r.font.size = Pt(size)
            self._shade(c, "D9E2F3")
        self._repeat_header(t.rows[0])
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                par = cells[i].paragraphs[0]
                par.paragraph_format.line_spacing = 1.0
                par.paragraph_format.space_after = Pt(0)
                lines = str(val).split("\n")
                for k, line in enumerate(lines):
                    if k:
                        par = cells[i].add_paragraph()
                        par.paragraph_format.line_spacing = 1.0
                        par.paragraph_format.space_after = Pt(0)
                    self._runs(par, line, size=size)
                if str(val) in ("Pass", "Met"):
                    self._shade(cells[i], "E2F0D9")
                elif str(val) in ("Fail", "Not met"):
                    self._shade(cells[i], "F8D7DA")
                elif str(val).startswith("Partly"):
                    self._shade(cells[i], "FFF2CC")
        for row in t.rows:
            for i, c in enumerate(row.cells):
                c.width = col_w[i]
                for par in c.paragraphs:
                    par.paragraph_format.line_spacing = 1.0
                    par.paragraph_format.space_after = Pt(0)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def _shade(self, cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    def _repeat_header(self, row):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trPr.append(el)

    def _b_fig(self, path, caption, width_cm, max_h_cm):
        self.fig_no += 1
        full = path if os.path.isabs(path) else os.path.join(FIG, path)
        with Image.open(full) as im:
            w, h = im.size
        uw, uh = self._usable()
        width = min(Cm(width_cm), uw)
        height = width * h / w
        limit = min(Cm(max_h_cm), uh - Cm(2))
        if height > limit:
            width = int(width * limit / height)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(full, width=int(width))
        self._caption("Figure", self.fig_no, caption, before=False)

    def _b_code(self, caption, file, start, end, lines):
        self.lst_no += 1
        code, a, b = extract(file, start, end, lines)
        self._caption("Listing", self.lst_no, f"{caption} (`{file}`, lines {a}–{b})".replace("`", ""), before=True)
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        cell = t.rows[0].cells[0]
        self._shade(cell, "F5F5F5")
        cell.text = ""
        for k, line in enumerate(code):
            par = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_after = Pt(0)
            r = par.add_run(line if line else " ")
            r.font.name = CODE_FONT
            r._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
            r.font.size = Pt(8.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def save(self, path):
        self.doc.save(path)
